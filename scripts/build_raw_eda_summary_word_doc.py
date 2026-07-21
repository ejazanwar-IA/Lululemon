#!/usr/bin/env python3
"""Build Word doc from findings/raw_sourcing_dev_7_table_eda_summary.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

PAGE_TABLE_WIDTH_DXA = 11483
PAGE_TABLE_INDENT_DXA = -1310

ROOT = Path(__file__).resolve().parents[1]
MD_IN = ROOT / "findings" / "raw_sourcing_dev_7_table_eda_summary.md"
DOC_OUT = ROOT / "findings" / "Lululemon_Raw_GCS_Extraction_EDA_Summary.docx"


def set_doc_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _column_widths_dxa(ncols: int) -> list[int]:
    total = PAGE_TABLE_WIDTH_DXA
    if ncols == 2:
        return [int(total * 0.31), total - int(total * 0.31)]
    if ncols == 3:
        a, b = int(total * 0.28), int(total * 0.22)
        return [a, b, total - a - b]
    if ncols == 4:
        a, b, c = int(total * 0.22), int(total * 0.18), int(total * 0.20)
        return [a, b, c, total - a - b - c]
    if ncols == 5:
        a, b, c, d = int(total * 0.12), int(total * 0.18), int(total * 0.30), int(total * 0.22)
        return [a, b, c, d, total - a - b - c - d]
    base = total // ncols
    widths = [base] * ncols
    widths[-1] += total - sum(widths)
    return widths


def style_full_width_table(table, ncols: int) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(PAGE_TABLE_WIDTH_DXA))
    tblW.set(qn("w:type"), "dxa")

    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(PAGE_TABLE_INDENT_DXA))
    tblInd.set(qn("w:type"), "dxa")

    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    widths = _column_widths_dxa(ncols)
    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is None:
        tblGrid = OxmlElement("w:tblGrid")
        tbl.insert(1, tblGrid)
    for gc in list(tblGrid.findall(qn("w:gridCol"))):
        tblGrid.remove(gc)
    for w in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(w))
        tblGrid.append(grid_col)

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def align_flush_left(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:left"), str(PAGE_TABLE_INDENT_DXA))


def clean_md_inline(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    return s


def add_paragraph_with_bold(doc: Document, text: str, *, bullet: bool = False) -> None:
    """Render **bold** segments in a paragraph."""
    text = text.strip()
    if not text:
        return
    style = "List Bullet" if bullet else None
    p = doc.add_paragraph(style=style)
    align_flush_left(p)
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(clean_md_inline(part))


def add_table_from_rows(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=ncols)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        padded = list(row) + [""] * ncols
        padded = padded[:ncols]
        for ci, val in enumerate(padded):
            cells[ci].text = clean_md_inline(str(val))
    style_full_width_table(table, ncols)
    doc.add_paragraph()


def add_title_block(doc: Document) -> None:
    t = doc.add_heading("Lululemon — Raw GCS Extraction EDA Summary", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in [
        "Client: Lululemon",
        "GCP project: lululemon-01042025",
        "Dataset: lulu_sourcing_dev",
        "Source: 7 GCS extraction tables (May 2026 EDAs)",
        "Last updated: 2026-06-23",
        "Prepared by: Impact Analytics",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_from_markdown(md_path: Path) -> Document:
    doc = Document()
    set_doc_styles(doc)
    add_title_block(doc)
    doc.add_page_break()

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("# "):
            h = doc.add_heading(clean_md_inline(line[2:]), 1)
            align_flush_left(h)
        elif line.startswith("## "):
            h = doc.add_heading(clean_md_inline(line[3:]), 2)
            align_flush_left(h)
        elif line.startswith("### "):
            h = doc.add_heading(clean_md_inline(line[4:]), 3)
            align_flush_left(h)
        elif line.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\|[-:\s|]+\|$", lines[i + 1]
        ):
            headers = [clean_md_inline(c.strip()) for c in line.strip("|").split("|")]
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(
                    [clean_md_inline(c.strip()) for c in lines[i].strip("|").split("|")]
                )
                i += 1
            add_table_from_rows(doc, headers, rows)
            continue
        elif line.strip().startswith("- "):
            add_paragraph_with_bold(doc, line.strip()[2:], bullet=True)
        elif line.strip() in ("---", ""):
            pass
        elif line.strip().startswith("`Last updated"):
            p = doc.add_paragraph(clean_md_inline(line.strip().strip("`")))
            align_flush_left(p)
        elif line.strip():
            add_paragraph_with_bold(doc, line)
        i += 1

    return doc


def main() -> None:
    if not MD_IN.exists():
        raise SystemExit(f"Missing input: {MD_IN}")
    doc = build_from_markdown(MD_IN)
    doc.save(DOC_OUT)
    print(f"Wrote {DOC_OUT} ({DOC_OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
